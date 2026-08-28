"""Build a self-contained DOCX with embedded DWG symbol example images."""

from __future__ import annotations

import argparse
from pathlib import Path
import tempfile

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_font(style, name: str, size: float | None = None) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        style.font.size = Pt(size)


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    _set_font(document.styles["Normal"], "Arial", 10.5)
    _set_font(document.styles["Title"], "Arial", 20)
    _set_font(document.styles["Heading 1"], "Arial", 15)
    _set_font(document.styles["Heading 2"], "Arial", 12)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)


def _render_svg(source: Path, destination: Path) -> None:
    drawing = pymupdf.open(stream=source.read_bytes(), filetype="svg")
    try:
        pixmap = drawing[0].get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=False)
        pixmap.save(destination)
    finally:
        drawing.close()


def _picture(
    document: Document,
    review_root: Path,
    relative_path: str,
    caption: str,
    temporary: Path,
) -> None:
    source = review_root / relative_path
    if not source.is_file():
        document.add_paragraph(f"Изображение отсутствует: {relative_path}")
        return
    png = temporary / f"image-{len(list(temporary.glob('*.png'))):03d}.png"
    _render_svg(source, png)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(png), width=Cm(14.5))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True


def _bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def _field(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value)


def build_report(review_root: Path, destination: Path) -> None:
    document = Document()
    _configure(document)
    document.core_properties.title = (
        "Краткий отчёт по распознаванию обозначений в DWG"
    )
    document.core_properties.subject = (
        "Примеры probable и нераспознанных обозначений"
    )

    title = document.add_heading(
        "Краткий отчёт по распознаванию обозначений в DWG",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Дата проверки: 26.08.2026\n").bold = True
    intro.add_run(
        "10 DWG-файлов, по одной выбранной странице каждого файла · этапы H2–H4c"
    )

    document.add_heading("Ограничение текущей реализации", level=1)
    document.add_paragraph(
        "В настоящее время легенда привязана к конкретному листу. Легенда "
        "извлекается с анализируемой страницы, и значки сопоставляются только "
        "с легендой этой же страницы."
    )
    _bullet(document, "Легенда с другого листа автоматически не применяется.")
    _bullet(document, "Легенда из другого DWG-файла комплекта не применяется.")
    _bullet(
        document,
        "Отсутствие локальной легенды не доказывает отсутствие обозначений.",
    )
    _bullet(
        document,
        "Confirmed и probable относятся только к локальной связи "
        "«значок — легенда на том же листе».",
    )

    document.add_heading("Общий результат", level=1)
    _bullet(document, "46 извлечённых пунктов локальных легенд.")
    _bullet(document, "82 совпадения confirmed.")
    _bullet(document, "48 совпадений probable.")
    _bullet(document, "338 нераспознанных кластеров: 1338 экземпляров.")
    _bullet(document, "Сопоставления получены на 5 из 10 проверенных листов.")
    document.add_paragraph(
        "Нераспознанные кандидаты включают не только условные обозначения, "
        "но также оборудование, двери, колонны, мебель, штампы и служебные блоки."
    )

    with tempfile.TemporaryDirectory(prefix="dwg_docx_") as directory:
        temporary = Path(directory)

        document.add_heading("Примеры probable", level=1)
        document.add_heading(
            "1. Перегородка из гипсокартонных листов",
            level=2,
        )
        _field(
            document,
            "Файл",
            "6 - ТХ/3. Жуковский 1_Блок 1-7 (ТХ)_Планы.dwg",
        )
        _field(document, "Страница", "1")
        _field(document, "Статус", "probable, confidence 0.85")
        _picture(
            document,
            review_root,
            (
                "tx-block1-plans-p1/dwg_symbols/page_0001/legend_crops/"
                "LE-ee1cc8751b98cda8.svg"
            ),
            "Образец из локальной легенды.",
            temporary,
        )
        _picture(
            document,
            review_root,
            (
                "tx-block1-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                "recognized-SI-2385829fffbc5361.svg"
            ),
            "Найденный на чертеже геометрический профиль.",
            temporary,
        )
        document.add_paragraph(
            "Геометрический стиль совпал с образцом легенды после нормализации "
            "поворота и длины. Результат не считается confirmed, поскольку он "
            "основан на стиле геометрии, а не на идентичном DWG-блоке."
        )

        document.add_heading(
            "2. Перегородка на архитектурном плане блока 8",
            level=2,
        )
        _field(
            document,
            "Файл",
            (
                "4 - КР/Раздел ПД №4 Часть 2 (КР2)/Внешние ссылки/"
                "Жуковский 1_блок 8_Планы.dwg"
            ),
        )
        _field(document, "Страница", "1")
        _field(document, "Статус", "probable, confidence 0.85")
        _picture(
            document,
            review_root,
            (
                "ar-block8-plans-p1/dwg_symbols/page_0001/legend_crops/"
                "LE-b75ce8aa0a180fdc.svg"
            ),
            "Локальный образец: перегородка из гипсокартонных листов.",
            temporary,
        )
        _picture(
            document,
            review_root,
            (
                "ar-block8-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                "recognized-SI-695da97aa3574e7f.svg"
            ),
            "Probable-экземпляр на архитектурном плане блока 8.",
            temporary,
        )
        document.add_paragraph(
            "Найдено совпадение нормализованного геометрического стиля с "
            "локальным образцом легенды. Длина участка — около 29 мм в "
            "координатах листа."
        )

        document.add_heading("3. Контекстное сопоставление по слою", level=2)
        _field(
            document,
            "Файл",
            "6 - ТХ/3. Жуковский 1_Блок 1-7 (ТХ)_Планы.dwg",
        )
        _field(document, "Страница", "1")
        _field(document, "Статус", "probable, confidence 0.80")
        _picture(
            document,
            review_root,
            (
                "tx-block1-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                "recognized-SI-1c0766b134bc3e01.svg"
            ),
            "Экземпляр, сопоставленный по контексту слоя БЛОКИ.",
            temporary,
        )
        document.add_paragraph(
            "Предполагаемое описание: стена или перегородка из блоков "
            "ячеистого бетона по ГОСТ 31360-2007. Имя слоя согласуется с "
            "подписью легенды, но это слабое доказательство и требует проверки."
        )

        document.add_heading("Примеры нераспознанных кандидатов", level=1)
        examples = [
            (
                "1. Блок трап100",
                "Технология",
                "трап100",
                "1",
                (
                    "tx-block1-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                    "unknown-SI-5c55c9d647f7da82.svg"
                ),
                (
                    "Блок присутствует на чертеже, но на листе не найден его "
                    "точный образец в локальной легенде. Название позволяет "
                    "предположить назначение, но статус не подтверждён."
                ),
            ),
            (
                "2. Сантехнический объект",
                "АР_сантех",
                "M_BATH_BASIN_Basin - Rect_P",
                "3",
                (
                    "tx-block1-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                    "unknown-SI-13a2d3f723647793.svg"
                ),
                (
                    "Имя блока похоже на сантехнический объект, однако "
                    "соответствующий пункт локальной легенды не найден."
                ),
            ),
            (
                "3. Повторяющийся блок колонны",
                "колонна",
                "Колонна",
                "54",
                (
                    "tx-block1-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                    "unknown-SI-61dd173345d4e62c.svg"
                ),
                (
                    "Unrecognized не обязательно означает неизвестный знак. "
                    "Это понятный конструктивный объект, отсутствующий в "
                    "локальной легенде данного листа."
                ),
            ),
            (
                "4. Служебный блок штампа",
                "FORMAT",
                "*U704",
                "1",
                (
                    "tx-block1-plans-p1/dwg_symbols/page_0001/symbol_crops/"
                    "unknown-SI-85da53abb33bb103.svg"
                ),
                (
                    "Элемент содержит номер листа, стадию, формат и должности. "
                    "Это оформление, а не условное обозначение, что показывает "
                    "необходимость отдельной фильтрации служебных блоков."
                ),
            ),
        ]
        for heading, layer, block, occurrences, image, description in examples:
            document.add_heading(heading, level=2)
            _field(
                document,
                "Файл",
                "6 - ТХ/3. Жуковский 1_Блок 1-7 (ТХ)_Планы.dwg",
            )
            _field(document, "Страница", "1")
            _field(document, "Слой", layer)
            _field(document, "DWG-блок", block)
            _field(document, "Экземпляров в кластере", occurrences)
            _picture(document, review_root, image, heading, temporary)
            document.add_paragraph(description)

    document.add_heading("Листы без применимой локальной легенды", level=1)
    _bullet(
        document,
        "ТХ, блок 8: найдено 5 текстовых пунктов, но нет пригодных "
        "векторных образцов; после fail-closed проверки — 0 сопоставлений.",
    )
    _bullet(document, "ТХ, блок 9: легенда на выбранной странице не извлечена.")
    _bullet(document, "ТХ, блок 10: легенда на выбранной странице не извлечена.")
    _bullet(document, "КР1, фундаменты: локальная легенда не найдена.")
    _bullet(document, "ПЗУ: результат неполный из-за отсутствующих XREF.")
    document.add_paragraph(
        "Нулевое число сопоставлений означает недостаток локальных "
        "доказательств, а не доказанное отсутствие условных обозначений."
    )

    document.add_heading("Вывод и следующий этап", level=1)
    document.add_paragraph(
        "Текущий подход работает, когда на анализируемом листе есть "
        "извлекаемая легенда. Следующий этап — сформировать общую легенду "
        "уровня DWG-файла и проектного комплекта."
    )
    steps = [
        "Объединить легенды всех листов одного файла.",
        "Разрешить наследование легенды между связанными файлами комплекта.",
        "Сохранять файл и лист, откуда взят образец.",
        "Отделить штампы и служебные блоки от символов чертежа.",
        "Подключить ручную разметку и внешние ГОСТ-каталоги для остатка.",
    ]
    for item in steps:
        document.add_paragraph(item, style="List Number")

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--review-root",
        type=Path,
        default=Path("local_runs/review_10"),
    )
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    destination = args.out or (
        args.review_root / "КРАТКИЙ_ОТЧЕТ_С_ПРИМЕРАМИ.docx"
    )
    build_report(args.review_root, destination)
    print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
